import io
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem


def export_to_docx(brd: dict) -> io.BytesIO:
    doc = Document()
    doc.add_heading(brd.get("title", "Business Requirements Document"), level=0)

    doc.add_heading("Overview", level=1)
    doc.add_paragraph(brd.get("overview", ""))

    doc.add_heading("Stakeholders", level=1)
    for s in brd.get("stakeholders", []):
        doc.add_paragraph(s, style="List Bullet")

    doc.add_heading("Functional Requirements", level=1)
    for fr in brd.get("functional_requirements", []):
        doc.add_paragraph(f"{fr.get('id', '')}: {fr.get('description', '')}", style="List Bullet")

    doc.add_heading("Non-Functional Requirements", level=1)
    for nfr in brd.get("non_functional_requirements", []):
        doc.add_paragraph(nfr, style="List Bullet")

    doc.add_heading("User Stories", level=1)
    for us in brd.get("user_stories", []):
        doc.add_paragraph(
            f"As a {us.get('as_a', '')}, I want {us.get('i_want', '')}, so that {us.get('so_that', '')}.",
            style="List Bullet"
        )

    doc.add_heading("Acceptance Criteria", level=1)
    for ac in brd.get("acceptance_criteria", []):
        doc.add_paragraph(ac, style="List Bullet")

    doc.add_heading("Assumptions", level=1)
    for a in brd.get("assumptions", []):
        doc.add_paragraph(a, style="List Bullet")

    doc.add_heading("Open Questions / Ambiguities", level=1)
    for q in brd.get("open_questions", []):
        doc.add_paragraph(q, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def export_to_pdf(brd: dict) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(brd.get("title", "Business Requirements Document"), styles["Title"]))
    elements.append(Spacer(1, 12))

    def add_section(heading, items, is_list=True):
        elements.append(Paragraph(heading, styles["Heading2"]))
        if is_list and items:
            bullets = [ListItem(Paragraph(str(i), styles["Normal"])) for i in items]
            elements.append(ListFlowable(bullets, bulletType="bullet"))
        elif not is_list:
            elements.append(Paragraph(str(items), styles["Normal"]))
        elements.append(Spacer(1, 10))

    add_section("Overview", brd.get("overview", ""), is_list=False)
    add_section("Stakeholders", brd.get("stakeholders", []))
    add_section(
        "Functional Requirements",
        [f"{fr.get('id', '')}: {fr.get('description', '')}" for fr in brd.get("functional_requirements", [])]
    )
    add_section("Non-Functional Requirements", brd.get("non_functional_requirements", []))
    add_section(
        "User Stories",
        [f"As a {us.get('as_a', '')}, I want {us.get('i_want', '')}, so that {us.get('so_that', '')}."
         for us in brd.get("user_stories", [])]
    )
    add_section("Acceptance Criteria", brd.get("acceptance_criteria", []))
    add_section("Assumptions", brd.get("assumptions", []))
    add_section("Open Questions / Ambiguities", brd.get("open_questions", []))

    doc.build(elements)
    buffer.seek(0)
    return buffer